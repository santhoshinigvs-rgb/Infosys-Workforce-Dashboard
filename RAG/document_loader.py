from pathlib import Path
import json
import pandas as pd
from docx import Document
from openpyxl import load_workbook

from rag.config import DATA_DIR, ML_DIR


SUPPORTED_TEXT = {".txt", ".md"}
SUPPORTED_TABLES = {".csv", ".xlsx", ".xls"}
SUPPORTED_DOCS = {".docx"}


def _chunk(text, source, chunk_size=1200, overlap=200):
    text = " ".join(str(text).split())
    if not text:
        return []

    chunks = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        chunks.append({
            "source": source,
            "text": text[start:end]
        })
        if end == len(text):
            break
        start = max(end - overlap, start + 1)
    return chunks


def load_file(path):
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix in SUPPORTED_TEXT:
        return _chunk(path.read_text(encoding="utf-8", errors="ignore"), str(path))

    if suffix == ".docx":
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return _chunk(text, str(path))

    if suffix == ".csv":
        df = pd.read_csv(path)
        text = (
            f"File: {path.name}\n"
            f"Columns: {', '.join(map(str, df.columns))}\n"
            f"Rows: {len(df)}\n"
            f"Sample records:\n{df.head(20).to_csv(index=False)}"
        )
        return _chunk(text, str(path))

    if suffix in {".xlsx", ".xls"}:
        if suffix == ".xlsx":
            wb = load_workbook(path, read_only=True, data_only=True)
            parts = [f"File: {path.name}"]
            for ws in wb.worksheets:
                rows = list(ws.iter_rows(values_only=True))
                parts.append(f"Sheet: {ws.title}")
                for row in rows[:30]:
                    parts.append(" | ".join("" if v is None else str(v) for v in row))
            return _chunk("\n".join(parts), str(path))

    return []


def discover_files():
    files = []

    for folder in (DATA_DIR, ML_DIR):
        if not folder.exists():
            continue

        for path in folder.rglob("*"):
            if path.is_file() and path.suffix.lower() in (
                SUPPORTED_TEXT | SUPPORTED_TABLES | SUPPORTED_DOCS
            ):
                files.append(path)

    return files


def load_all_documents():
    documents = []
    for path in discover_files():
        try:
            documents.extend(load_file(path))
        except Exception as exc:
            documents.append({
                "source": str(path),
                "text": f"Could not parse this file: {exc}"
            })
    return documents
